# -*- coding: utf-8 -*-
"""
feature_extractor.py
-------------------
Extracts, encodes, and standardizes pre-publication and target features
from the relational External Intelligence Benchmark Corpus (XLSX).

Enforces strict target leakage prevention:
- Pre-publication features (safe for inference): hook encodings, duration, lexical stats,
  code density, analogy indicators, channel size tiers.
- Primary target: computed_relative_views (log_relative_views for regression).
- Secondary target: computed_performance_class (for classification).
- Independent annotations (never features): researcher_relative_performance, researcher_performance_class.
"""

import re
import math
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
import openpyxl

TECH_TERMS = {
    'api', 'apis', 'algorithm', 'algorithms', 'array', 'arrays', 'async', 'await',
    'backend', 'frontend', 'database', 'sql', 'nosql', 'docker', 'kubernetes',
    'git', 'github', 'linux', 'terminal', 'bash', 'python', 'javascript',
    'typescript', 'rust', 'golang', 'c++', 'cssz', 'html', 'react', 'node',
    'server', 'client', 'http', 'rest', 'graphql', 'orm', 'ast', 'compiler',
    'runtime', 'memory', 'pointer', 'mutex', 'concurrency', 'thread', 'threads',
    'lambda', 'closure', 'hash', 'hashmap', 'encryption', 'hashing', 'security',
    'token', 'llm', 'llms', 'ai', 'transformer', 'cuda', 'gpu', 'vps', 'ssh',
    'tmux', 'vim', 'neovim', 'refactor', 'refactoring', 'linter', 'ci/cd'
}

CODE_TOKENS = {
    'def', 'function', 'const', 'let', 'var', 'import', 'from', 'return',
    'class', 'if', 'else', 'for', 'while', 'try', 'catch', 'except', 'lambda',
    'print', 'console.log', 'malloc', 'free', 'struct', 'curl', 'npm', 'pip',
    'docker-compose', 'sudo', 'grep', 'awk', 'sed'
}


def extract_lexical_features(transcript: Optional[str]) -> Dict[str, float]:
    """Programmatically derives safe lexical signals from verbatim transcripts."""
    if not transcript or not isinstance(transcript, str) or not transcript.strip():
        return {
            'transcript_char_count': 0,
            'transcript_word_count': 0,
            'transcript_sentence_count': 0,
            'question_count': 0,
            'question_density': 0.0,
            'number_count': 0,
            'number_density': 0.0,
            'tech_term_count': 0,
            'tech_term_density': 0.0,
            'code_token_count': 0,
            'code_token_density': 0.0,
            'avg_word_length': 0.0,
            'words_per_second': 0.0
        }

    text = transcript.strip()
    char_count = len(text)
    raw_words = re.findall(r'[a-zA-Z0-9_\-\+\#\.]+', text.lower())
    word_count = len(raw_words)
    sentences = [s.strip() for s in re.split(r'[.!?]+', text) if s.strip()]
    sentence_count = max(1, len(sentences))
    question_count = text.count('?')
    question_density = round(question_count / max(1, sentence_count), 3)
    numbers = re.findall(r'\b\d+(?:\.\d+)?%?\b', text)
    number_count = len(numbers)
    number_density = round(number_count / max(1, word_count), 4)
    tech_matches = [w for w in raw_words if w in TECH_TERMS]
    tech_term_count = len(tech_matches)
    tech_term_density = round(tech_term_count / max(1, word_count), 4)
    code_matches = [w for w in raw_words if w in CODE_TOKENS]
    code_token_count = len(code_matches)
    code_token_density = round(code_token_count / max(1, word_count), 4)
    total_word_chars = sum(len(w) for w in raw_words)
    avg_word_length = round(total_word_chars / max(1, word_count), 2)

    return {
        'transcript_char_count': char_count,
        'transcript_word_count': word_count,
        'transcript_sentence_count': sentence_count,
        'question_count': question_count,
        'question_density': question_density,
        'number_count': number_count,
        'number_density': number_density,
        'tech_term_count': tech_term_count,
        'tech_term_density': tech_term_density,
        'code_token_count': code_token_count,
        'code_token_density': code_token_density,
        'avg_word_length': avg_word_length,
        'words_per_second': 0.0
    }


class BenchmarkFeaturePipeline:
    """Extracts and processes tabular research sheets into an ML-ready feature matrix."""

    def __init__(self, workbook_path: Path):
        self.workbook_path = Path(workbook_path)
        if not self.workbook_path.exists():
            raise FileNotFoundError(f'Workbook not found: {self.workbook_path}')

    def load_relational_tables(self) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        wb = openpyxl.load_workbook(self.workbook_path, data_only=True)

        def sheet_to_rows(s_name: str) -> List[Dict[str, Any]]:
            if s_name not in wb.sheetnames:
                return []
            ws = wb[s_name]
            headers = [str(c.value).strip() if c.value is not None else f'col_{idx}' for idx, c in enumerate(ws[1])]
            rows = []
            for r_cells in ws.iter_rows(min_row=2, values_only=True):
                r_dict = {h: val for h, val in zip(headers, r_cells)}
                if any(v is not None and str(v).strip() != '' for v in r_dict.values()):
                    rows.append(r_dict)
            return rows

        channels_list = sheet_to_rows('CHANNELS')
        videos_list = sheet_to_rows('VIDEOS')
        transcripts_list = sheet_to_rows('TRANSCRIPTS')
        content_list = sheet_to_rows('CONTENT_ANALYSIS')
        perf_list = sheet_to_rows('PERFORMANCE')

        channels_map = {str(c['channel_id']): c for c in channels_list if c.get('channel_id')}
        transcripts_map = {str(t['video_id']): t for t in transcripts_list if t.get('video_id')}
        content_map = {str(cn['video_id']): cn for cn in content_list if cn.get('video_id')}
        perf_map = {str(p['video_id']): p for p in perf_list if p.get('video_id')}

        valid_rows: List[Dict[str, Any]] = []
        quarantined_rows: List[Dict[str, Any]] = []
        seen_vids = set()

        for v in videos_list:
            raw_vid = str(v.get('video_id') or '').strip()
            cid = str(v.get('channel_id') or '').strip()
            title = str(v.get('video_title') or '').strip()
            url = str(v.get('youtube_url') or '').strip()
            is_valid_id = bool(re.match(r'^[a-zA-Z0-9_-]{11}$', raw_vid))

            ch = channels_map.get(cid, {})
            tr = transcripts_map.get(raw_vid, {})
            co = content_map.get(raw_vid, {})
            pe = perf_map.get(raw_vid, {})

            record = {
                'video_id': raw_vid,
                'video_title': title,
                'youtube_url': url,
                'channel_id': cid,
                'channel_name': ch.get('channel_name', 'Unknown'),
                'channel_size_class': ch.get('channel_size_class', 'UNKNOWNE'),
                'public_subscriber_count': ch.get('public_subscriber_count'),
                'channel_total_video_count': ch.get('channel_total_video_count'),
                'published_at': v.get('published_at'),
                'duration_seconds': float(v.get('duration_seconds') or 45.0),
                'is_short': bool(v.get('is_short', True)),
                'video_topic': v.get('video_topic'),
                'video_language': v.get('video_language', 'en'),
                'transcript_status': tr.get('transcript_status', 'UNAVAILABLE'),
                'transcript': tr.get('transcript'),
                'hook_structure': co.get('hook_structure', 'UNKNOWNE'),
                'explanation_structure': co.get('explanation_structure', 'UNKNOWN'),
                'analogy_usage': bool(co.get('analogy_usage', False)),
                'code_usage': bool(co.get('code_usage', False)),
                'example_usage': bool(co.get('example_usage', False)),
                'visual_technique': co.get('visual_technique', 'UNKNOWN'),
                'cta_structure': co.get('cta_structure', 'UNKNOWN'),
                'technical_level': co.get('technical_level', 'Beginner'),
                'primary_topic': co.get('primary_topic', 'General'),
                'secondary_topic': co.get('secondary_topic', 'General'),
                'audience_problem': co.get('audience_problem'),
                'simplification_method': co.get('simplification_method', 'UNKNOWN'),
                'presentation_style': co.get('presentation_style', 'UNKNOWN'),
                'information_density': co.get('information_density', 'Moderate'),
                'views': int(pe.get('views') or 0),
                'likes': int(pe['likes']) if pe.get('likes') is not None else None,
                'comments': int(pe['comments']) if pe.get('comments') is not None else None,
                'channel_median_views': float(pe.get('channel_median_views') or 100000.0),
                'researcher_relative_performance': float(pe.get('researcher_relative_performance') or 1.0),
                'researcher_performance_class': pe.get('researcher_performance_class', 'UNCLASSIFIED'),
                'computed_relative_views': float(pe.get('computed_relative_views') or 1.0),
                'computed_performance_class': pe.get('computed_performance_class', 'NORMAL_PERFORMER')
            }

            if not is_valid_id:
                record['rejection_reason'] = 'Invalid YouTube video_id format (must be 11-char canonical ID)'
                quarantined_rows.append(record)
            elif raw_vid in seen_vids:
                record['rejection_reason'] = f"Duplicate video_id '{raw_vid}'"
                quarantined_rows.append(record)
            else:
                seen_vids.add(raw_vid)
                valid_rows.append(record)

        return valid_rows, quarantined_rows

    def build_ml_feature_matrix(self) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        valid_rows, quarantined_rows = self.load_relational_tables()
        ml_rows: List[Dict[str, Any]] = []

        for r in valid_rows:
            duration = max(5.0, float(r['duration_seconds']))
            lex = extract_lexical_features(r['transcript'])
            lex['words_per_second'] = round(lex['transcript_word_count'] / duration, 2)

            rel_views = max(0.01, float(r['computed_relative_views']))
            log_rel_views = round(math.log(rel_views), 4)

            hook = str(r['hook_structure']).upper()
            expl = str(r['explanation_structure']).upper()
            vis = str(r.get('visual_technique') or '').upper()
            tech_lvl = str(r['technical_level']).capitalize()
            density = str(r['information_density']).capitalize()
            size_cls = str(r['channel_size_class']).upper()
            cta = str(r['cta_structure']).upper()

            feature_dict = {
                'meta_video_id': r['video_id'],
                'meta_channel_id': r['channel_id'],
                'meta_channel_name': r['channel_name'],
                'meta_video_title': r['video_title'],
                'meta_published_at': r['published_at'],
                'meta_raw_views': r['views'],
                'meta_channel_median_views': r['channel_median_views'],

                'cat_hook_structure': r['hook_structure'],
                'cat_explanation_structure': r['explanation_structure'],
                'cat_visual_technique': r['visual_technique'],
                'cat_technical_level': r['technical_level'],
                'cat_information_density': r['information_density'],
                'cat_cta_structure': r['cta_structure'],

                'target_computed_relative_views': rel_views,
                'target_log_relative_views': log_rel_views,
                'target_computed_performance_class': r['computed_performance_class'],
                'target_is_high_performer': 1 if r['computed_performance_class'] == 'HIGH_PERFORMER' else 0,
                'target_is_low_performer': 1 if r['computed_performance_class'] == 'LOW_PERFORMER' else 0,

                'annot_researcher_relative_perf': r['researcher_relative_performance'],
                'annot_researcher_class': r['researcher_performance_class'],
                'annot_has_classification_disagreement': 1 if r['researcher_performance_class'] != r['computed_performance_class'] else 0,

                'feat_channel_size_large': 1 if size_cls == 'LARGE' else 0,
                'feat_channel_size_medium': 1 if size_cls == 'MEDIUM' else 0,
                'feat_channel_size_small': 1 if size_cls == 'SMALL' else 0,
                'feat_public_subscribers': r['public_subscriber_count'] or 0,
                'feat_log_subscribers': round(math.log(max(100, r['public_subscriber_count'] or 100)), 2),

                'feat_duration_seconds': duration,
                'feat_is_short_duration_under_50s': 1 if duration < 50.0 else 0,

                'feat_hook_counterintuitive': 1 if 'COUNTERINTUITIVE' in hook else 0,
                'feat_hook_question': 1 if 'QUESTION' in hook else 0,
                'feat_hook_mistake_first': 1 if 'MISTAKE' in hook else 0,
                'feat_hook_problem_first': 1 if 'PROBLEM' in hook else 0,
                'feat_hook_curiosity_gap': 1 if 'CURIOSITY' in hook else 0,
                'feat_hook_news_hook': 1 if 'NEWS' in hook else 0,
                'feat_hook_demonstration_first': 1 if 'DEMONSTRATION' in hook or 'RESULT' in hook else 0,

                'feat_expl_before_after': 1 if 'BEFORE_AFTER' in expl else 0,
                'feat_expl_mistake_why_fix': 1 if 'MISTAKE_WHY_FIX' in expl else 0,
                'feat_expl_concept_analogy_code': 1 if 'CONCEPT_ANALOGY_CODE' in expl else 0,
                'feat_expl_step_by_step': 1 if 'STEP_BY_STEP' in expl else 0,
                'feat_expl_demo_explanation': 1 if 'DEMO_EXPLANATION' in expl else 0,
                'feat_expl_news_context': 1 if 'NEWS_CONTEXT' in expl else 0,
                'feat_expl_story_explanation': 1 if 'STORY' in expl else 0,

                'feat_analogy_used': 1 if r.get('analogy_usage') else 0,
                'feat_code_used': 1 if r.get('code_usage') else 0,
                'feat_example_used': 1 if r.get('example_usage') else 0,

                'feat_visual_code_animation': 1 if 'CODE_ANIMATION' in vis or 'SPLIT_SCREEN' in vis else 0,
                'feat_visual_prop_or_hardware': 1 if 'PROP' in vis else 0,
                'feat_visual_screen_recording': 1 if 'SCREEN_RECORDING' in vis or 'IDE' in vis else 0,
                'feat_visual_high_energy_editing': 1 if 'HIGH_ENERGY' in vis else 0,

                'feat_cta_seamless_loop': 1 if 'SEAMLESS' in cta else 0,
                'feat_cta_verbal_prompt': 1 if 'VERBAL' in cta else 0,

                'feat_tech_level_beginner': 1 if tech_lvl == 'Beginner' else 0,
                'feat_tech_level_intermediate': 1 if tech_lvl == 'Intermediate' else 0,
                'feat_tech_level_advanced': 1 if tech_lvl == 'Advanced' else 0,

                'feat_density_high': 1 if density == 'High' else 0,
                'feat_density_moderate': 1 if density == 'Moderate' else 0,
                'feat_density_low': 1 if density == 'Low' else 0,

                'feat_transcript_word_count': lex['transcript_word_count'],
                'feat_words_per_second': lex['words_per_second'],
                'feat_question_density': lex['question_density'],
                'feat_number_density': lex['number_density'],
                'feat_tech_term_density': lex['tech_term_density'],
                'feat_code_token_density': lex['code_token_density'],
                'feat_avg_word_length': lex['avg_word_length']
            }
            ml_rows.append(feature_dict)

        return ml_rows, quarantined_rows
