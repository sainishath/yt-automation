# -*- coding: utf-8 -*-
"""
dataset_importer.py
-------------------
Master importer for tabular external research datasets (CSV, TSV, XLSX).
Performs file parsing, alias normalization, strict row-by-row validation,
channel-relative metric calculation, disk transcript persistence, and idempotent SQLite storage.
"""

import csv
import json
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime

from growth.db.database import get_db, DEFAULT_DB_PATH
from growth.external_intelligence.schemas import (
    ExternalChannelModel,
    ExternalVideoModel,
    ExternalVideoSnapshotModel,
    ExternalObservationModel,
    ObservationType,
    EvidenceLevel,
    ProvenanceSource
)
from growth.external_intelligence.repository import ExternalIntelligenceRepository
from growth.external_intelligence.ingestion.ingestion_schemas import (
    TranscriptStatus,
    PerformanceClass,
    ExternalDatasetRegistryModel,
    ValidatedExternalVideoRecord,
    IngestionReport
)
from growth.external_intelligence.ingestion.raw_validator import validate_record
from growth.external_intelligence.ingestion.normalizer import normalize_dataset_records

TRANSCRIPT_BASE_DIR = Path(__file__).parent.parent.parent.parent / "research" / "youtube_competitors" / "transcripts"


class TabularDatasetImporter:
    """
    Ingests, validates, and registers external competitor research datasets.
    Supports CSV, TSV, and XLSX formats with strict provenance and dry-run safety.
    """

    def __init__(
        self,
        repo: Optional[ExternalIntelligenceRepository] = None,
        db_path: Path = DEFAULT_DB_PATH,
        transcript_dir: Path = TRANSCRIPT_BASE_DIR
    ):
        self.db_path = db_path
        self.repo = repo or ExternalIntelligenceRepository(db_path)
        self.transcript_dir = transcript_dir
        self.transcript_dir.mkdir(parents=True, exist_ok=True)

    def parse_tabular_file(self, file_path: Path) -> List[Dict[str, Any]]:
        """Parses CSV, TSV, or XLSX file into a list of row dictionaries."""
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Dataset file not found: {file_path}")

        suffix = file_path.suffix.lower()
        rows: List[Dict[str, Any]] = []

        if suffix in (".csv", ".txt"):
            with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f:
                reader = csv.DictReader(f)
                for r in reader:
                    rows.append(dict(r))
        elif suffix == ".tsv":
            with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f:
                reader = csv.DictReader(f, delimiter="\t")
                for r in reader:
                    rows.append(dict(r))
        elif suffix == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                
                # Check if multi-sheet relational structure exists
                sheet_names_upper = {s.upper(): s for s in wb.sheetnames}
                if "VIDEOS" in sheet_names_upper:
                    # Multi-sheet relational workbook
                    def sheet_to_dicts(s_name: str, key_field: str = "video_id") -> Dict[str, Dict[str, Any]]:
                        if s_name not in sheet_names_upper:
                            return {}
                        ws = wb[sheet_names_upper[s_name]]
                        headers = [str(cell.value).strip() if cell.value is not None else f"col_{idx}" for idx, cell in enumerate(ws[1])]
                        mapping = {}
                        for r_cells in ws.iter_rows(min_row=2, values_only=True):
                            r_dict = {h: val for h, val in zip(headers, r_cells)}
                            k = str(r_dict.get(key_field) or "").strip()
                            if k:
                                mapping[k] = r_dict
                        return mapping

                    channels_map = sheet_to_dicts("CHANNELS", "channel_id")
                    videos_map = sheet_to_dicts("VIDEOS", "video_id")
                    transcripts_map = sheet_to_dicts("TRANSCRIPTS", "video_id")
                    content_map = sheet_to_dicts("CONTENT_ANALYSIS", "video_id")
                    perf_map = sheet_to_dicts("PERFORMANCE", "video_id")

                    for vid, v_row in videos_map.items():
                        merged = dict(v_row)
                        cid = str(v_row.get("channel_id") or "").strip()
                        if cid in channels_map:
                            for ch_k, ch_v in channels_map[cid].items():
                                if ch_k not in merged or merged[ch_k] is None:
                                    merged[ch_k] = ch_v
                        if vid in transcripts_map:
                            for t_k, t_v in transcripts_map[vid].items():
                                if t_k != "video_id":
                                    merged[t_k] = t_v
                        if vid in content_map:
                            for c_k, c_v in content_map[vid].items():
                                if c_k != "video_id":
                                    merged[c_k] = c_v
                        if vid in perf_map:
                            for p_k, p_v in perf_map[vid].items():
                                if p_k != "video_id":
                                    merged[p_k] = p_v
                        rows.append(merged)
                else:
                    # Single flat sheet workbook
                    sheet = wb.active
                    headers = [str(cell.value).strip() if cell.value is not None else f"col_{idx}" for idx, cell in enumerate(sheet[1])]
                    for row_cells in sheet.iter_rows(min_row=2, values_only=True):
                        row_dict = {}
                        for h, val in zip(headers, row_cells):
                            row_dict[h] = val
                        if any(v is not None and str(v).strip() != "" for v in row_dict.values()):
                            rows.append(row_dict)
            except ImportError:
                # Fallback message if openpyxl is not installed
                raise RuntimeError("openpyxl is required to parse .xlsx files. Please install openpyxl or supply .csv/.tsv format.")
        else:
            raise ValueError(f"Unsupported file extension '{suffix}'. Supported: .csv, .tsv, .xlsx")

        return rows

    def import_dataset(
        self,
        file_path: Path,
        dataset_id: Optional[str] = None,
        dataset_version: str = "v1.0",
        target_channel_id: str = "channel_b",
        dry_run: bool = False
    ) -> IngestionReport:
        """
        Executes end-to-end dataset import with validation, normalization, and persistence.
        """
        file_path = Path(file_path)
        if not dataset_id:
            stem = file_path.stem.lower().replace(" ", "_").replace("-", "_")
            dataset_id = f"ds_{stem}"

        raw_rows = self.parse_tabular_file(file_path)
        report = IngestionReport(
            dataset_id=dataset_id,
            file_path=str(file_path),
            total_rows=len(raw_rows),
            is_dry_run=dry_run
        )

        records: List[ValidatedExternalVideoRecord] = []
        seen_video_ids = set()

        # Step 1: Validate individual rows
        for idx, row in enumerate(raw_rows, start=1):
            rec = validate_record(idx, row, dataset_id)

            if rec.video_id in seen_video_ids and rec.is_valid:
                rec.is_valid = False
                rec.validation_errors.append(f"Duplicate video_id '{rec.video_id}' within same dataset batch")
                report.duplicate_rows += 1

            if rec.is_valid:
                seen_video_ids.add(rec.video_id)
                report.valid_rows += 1
                if rec.transcript_status == TranscriptStatus.UNAVAILABLE or not rec.transcript_text:
                    report.missing_transcripts += 1
                if rec.likes is None or rec.comments is None:
                    report.missing_metrics += 1
            else:
                report.rejected_rows += 1
                for err in rec.validation_errors:
                    report.rejection_reasons[err] = report.rejection_reasons.get(err, 0) + 1
                report.quarantine_log.append({
                    "row_index": idx,
                    "title": rec.video_title,
                    "video_id": rec.video_id,
                    "errors": rec.validation_errors
                })

            records.append(rec)

        # Step 2: Compute channel-relative normalization across valid rows
        normalized_records = normalize_dataset_records(records)

        # Build performance distribution
        for r in normalized_records:
            if r.is_valid:
                p_cls = r.computed_classification.value
                report.performance_distribution[p_cls] = report.performance_distribution.get(p_cls, 0) + 1

        # Step 3: Database & disk persistence (if not dry-run)
        if not dry_run:
            self._persist_dataset(
                dataset_id=dataset_id,
                dataset_version=dataset_version,
                target_channel_id=target_channel_id,
                file_path=file_path,
                report=report,
                records=normalized_records
            )

        return report

    def _persist_dataset(
        self,
        dataset_id: str,
        dataset_version: str,
        target_channel_id: str,
        file_path: Path,
        report: IngestionReport,
        records: List[ValidatedExternalVideoRecord]
    ) -> None:
        """Persists validated external records, channels, snapshots, and transcript files."""
        now_str = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

        # 1. Register or update dataset metadata in SQLite
        with get_db(self.db_path) as conn:
            conn.execute("""
                CREATE TABLE IF NOT EXISTS external_datasets (
                    dataset_id TEXT PRIMARY KEY,
                    dataset_version TEXT NOT NULL,
                    source TEXT NOT NULL,
                    collector TEXT NOT NULL,
                    schema_version TEXT NOT NULL,
                    row_count INTEGER DEFAULT 0,
                    valid_count INTEGER DEFAULT 0,
                    rejected_count INTEGER DEFAULT 0,
                    quarantined_count INTEGER DEFAULT 0,
                    file_path TEXT,
                    imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
            conn.execute("""
                INSERT OR REPLACE INTO external_datasets (
                    dataset_id, dataset_version, source, collector, schema_version,
                    row_count, valid_count, rejected_count, quarantined_count, file_path, imported_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                dataset_id,
                dataset_version,
                "GEMINI_PUBLIC_RESEARCH",
                "HUMAN_OPERATOR",
                "external_tabular_v1",
                report.total_rows,
                report.valid_rows,
                report.rejected_rows,
                len(report.quarantine_log),
                str(file_path),
                now_str
            ))

        # 2. Persist channels and videos
        for rec in records:
            if not rec.is_valid:
                continue

            # Save transcript to disk if available
            t_path = None
            if rec.transcript_text and rec.transcript_status != TranscriptStatus.UNAVAILABLE:
                transcript_file = self.transcript_dir / f"{rec.video_id}.txt"
                try:
                    transcript_file.write_text(rec.transcript_text, encoding="utf-8")
                    t_path = str(transcript_file)
                except Exception as e:
                    logging.warning(f"Could not write transcript file for {rec.video_id}: {e}")

            # Upsert External Channel
            channel_id = f"ext_ch_{rec.channel_name.lower().replace(' ', '_')[:30]}"
            ch_model = ExternalChannelModel(
                external_channel_id=channel_id,
                target_channel_id=target_channel_id,
                channel_title=rec.channel_name,
                handle=rec.channel_url,
                youtube_channel_id=rec.channel_url,
                subscriber_count=rec.public_subscriber_count or 0,
                content_niche="coding_tech_education",
                similarity_score=0.92,
                similarity_reasons=["Public technical education analog"],
                source_type=ProvenanceSource.PUBLIC_YOUTUBE,
                last_researched_at=now_str
            )
            self.repo.upsert_external_channel(ch_model)

            # Upsert External Video
            vid_model = ExternalVideoModel(
                external_video_id=f"ext_vid_{rec.video_id}",
                external_channel_id=channel_id,
                youtube_video_id=rec.video_id,
                title=rec.video_title,
                url=rec.youtube_url,
                published_at=rec.published_at,
                duration_seconds=rec.duration_seconds,
                is_short=rec.duration_seconds <= 90.0,
                views=rec.views,
                likes=rec.likes or 0,
                comments=rec.comments or 0,
                relative_view_multiplier=rec.computed_relative_views,
                is_simulation=False,
                source_type=ProvenanceSource.PUBLIC_YOUTUBE
            )
            self.repo.upsert_external_video(vid_model)

            # Insert initial snapshot
            snap_model = ExternalVideoSnapshotModel(
                external_video_id=f"ext_vid_{rec.video_id}",
                window_name="initial",
                views=rec.views,
                likes=rec.likes or 0,
                comments=rec.comments or 0,
                relative_view_multiplier=rec.computed_relative_views,
                source_type=ProvenanceSource.PUBLIC_YOUTUBE,
                is_simulation=False
            )
            self.repo.upsert_external_video_snapshot(snap_model)

            # Record detailed observations
            self.repo.insert_observation(ExternalObservationModel(
                observation_id=f"obs_perf_{rec.video_id}",
                external_video_id=f"ext_vid_{rec.video_id}",
                observation_type=ObservationType.OBJECTIVE_FACT,
                field_name="computed_classification",
                observed_value=rec.computed_classification.value,
                interpretation=f"Researcher claimed: {rec.researcher_classification.value} (Relative: {rec.computed_relative_views}x)",
                evidence_level=EvidenceLevel.LEVEL_2_EXTERNAL_EVIDENCE,
                source_type=ProvenanceSource.PUBLIC_YOUTUBE
            ))

            if rec.transcript_status != TranscriptStatus.UNAVAILABLE:
                self.repo.insert_observation(ExternalObservationModel(
                    observation_id=f"obs_transcript_{rec.video_id}",
                    external_video_id=f"ext_vid_{rec.video_id}",
                    observation_type=ObservationType.OBJECTIVE_FACT,
                    field_name="transcript_status",
                    observed_value=rec.transcript_status.value,
                    interpretation=f"Transcript path: {t_path}",
                    evidence_level=EvidenceLevel.LEVEL_1_OBSERVATION,
                    source_type=ProvenanceSource.PUBLIC_YOUTUBE
                ))
